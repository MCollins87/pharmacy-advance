#!/usr/bin/env python3
"""Build the four-tab ARIA pharmacy planning workbook.

Source precedence:
1. sch_inst_by_pt_sum_aal.xls is the authoritative daily schedule.
2. pharm_reqmt.xls is the preferred medication source. Matching lines are Approved.
3. ptmeds_sch_time.pdf is the fallback medication source. Matching lines are Planned.

The output is a planning/reconciliation aid. It does not calculate doses and it
must not replace verification against the current authorised ARIA prescription.
"""
from __future__ import annotations

import logging
import re
import shutil
import sys
from collections import defaultdict
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path

import xlrd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo
from pypdf import PdfReader

BASE_DIR = Path(r"C:\IDR\PharmacyAdvance")
INPUT_DIR = BASE_DIR / "Input"
OUTPUT_DIR = BASE_DIR / "Output"
ARCHIVE_DIR = BASE_DIR / "Archive"
LOG_DIR = BASE_DIR / "Logs"

SCHEDULE_FILENAME = "sch_inst_by_pt_sum_aal.xls"
PHARMACY_FILENAME = "pharm_reqmt.xls"
PDF_FILENAME = "ptmeds_sch_time.pdf"

SCHEDULE_REPORT_TITLE = "Schedule - Institution by Patient and Time - Summary"
PHARMACY_REPORT_TITLE = (
    "Pharmacy Requirements - by Agent, Rx Type, Administration Date and Patient"
)
PDF_REPORT_TITLE = "Patient Medications - Patients Scheduled to Visit - by Time"

OUTPUT_HEADERS = [
    "Administration date", "Appointment time", "Patient ID", "Patient",
    "Prescription physician", "Visit provider", "Scheduled event(s)",
    "Drug", "Dose", "Route", "Dispensed", "Verified", "Source",
    "Source reference", "Status / reason",
]

PDF_PATIENT_RE = re.compile(
    r"Start Time:\s*(?P<date>[A-Z][a-z]{2}\s+\d{1,2},\s+\d{4})\s+"
    r"(?P<time>\d{1,2}:\d{2})\s+Event Provider:.*?\n"
    r"(?P<patient>.+?)\s+NHS Number:\s*(?P<nhs>\d{10})\b", re.S,
)
PDF_SECTION_RE = re.compile(
    r"^(Chemo|Hormone|Immunotherapy|Supportive|Pharmacy:)\s*$", re.I
)
PDF_FOOTER_RE = re.compile(r"^Report Name:", re.I)
PDF_ADMIN_RE = re.compile(r"^\s*Administration Instructions:", re.I)
PDF_DATE_AT_END_RE = re.compile(
    r"\s+[A-Z][a-z]{2}\s+\d{1,2},\s+\d{4}"
    r"(?:\s+[A-Z][a-z]{2}\s+\d{1,2},\s+\d{4})?\s*$"
)


@dataclass(frozen=True)
class ScheduleRow:
    nhs_number: str
    patient: str
    event_dt: datetime
    event: str
    visit_provider: str
    location: str
    comments: str
    source_row: int


@dataclass(frozen=True)
class PharmacyRequirement:
    administration_date: date
    agent: str
    dose: str
    patient: str
    physician: str
    dispensed: str
    verified: str
    source_row: int


@dataclass(frozen=True)
class PdfPatient:
    visit_date: date
    visit_time: datetime
    nhs_number: str
    patient: str
    page: int


@dataclass(frozen=True)
class PdfMedication:
    visit_date: date
    visit_time: datetime
    nhs_number: str
    patient: str
    section: str
    agent: str
    course_description: str
    dose: str
    route: str
    page: int


def clean(value) -> str:
    return re.sub(r"\s+", " ", str(value or "")).strip()


def normalise(value) -> str:
    return re.sub(r"[^a-z0-9]+", "", clean(value).casefold())


def normalise_nhs_number(value) -> str:
    text = clean(value)
    if text.casefold() == "reserved":
        return "Reserved"
    digits = re.sub(r"\D", "", text)
    return digits or text


def patient_match_key(value) -> str:
    text = clean(value).upper()
    if not text:
        return ""
    if "," in text:
        surname, forenames = text.split(",", 1)
    else:
        parts = text.split()
        surname, forenames = parts[0], " ".join(parts[1:])
    names = re.findall(r"[A-Z0-9]+", forenames)
    return f"{normalise(surname)}|{normalise(names[0] if names else '')}"


def parse_excel_datetime(value, datemode: int) -> datetime | None:
    if value in (None, ""):
        return None
    if isinstance(value, (int, float)):
        return xlrd.xldate_as_datetime(value, datemode)
    text = clean(value)
    for fmt in ("%b %d, %Y %H:%M", "%d/%m/%Y %H:%M", "%Y-%m-%d %H:%M",
                "%b %d, %Y", "%d/%m/%Y", "%Y-%m-%d"):
        try:
            return datetime.strptime(text, fmt)
        except ValueError:
            pass
    return None


def parse_text_date(value) -> date | None:
    text = clean(value)
    for fmt in ("%b %d, %Y", "%d/%m/%Y", "%Y-%m-%d"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            pass
    return None


def format_dose(amount, unit) -> str:
    if amount in (None, ""):
        return ""
    if isinstance(amount, float) and amount.is_integer():
        amount = int(amount)
    return clean(f"{amount} {unit}")


def extract_pdf_dose(text: str) -> str:
    text = re.sub(r"^Prescribed:\s*", "", clean(text), flags=re.I)
    match = re.match(
        r"(?P<dose>\d[\d,]*(?:\.\d+)?(?:\s*-\s*\d[\d,]*(?:\.\d+)?)?\s*"
        r"(?:mg|mcg|g|mmol|mL|IU|Units|dose\(s\)|tablet|capsule|injection)"
        r"(?:\s*\([^)]*\))?)\b", text, re.I,
    )
    return clean(match.group("dose")) if match else ""


def extract_route(text: str) -> str:
    for route in ("Subcutaneous", "Intravenous", "Intramuscular", "Oral",
                  "Oromucosal", "Topical", "Ocular", "Inhalation", "Neb"):
        if re.search(rf"\b{re.escape(route)}\b", text, re.I):
            return route
    return ""


def configure_logging() -> Path:
    LOG_DIR.mkdir(parents=True, exist_ok=True)
    path = LOG_DIR / f"pharmacy_advance_{datetime.now():%Y%m%d_%H%M%S}.log"
    logging.basicConfig(
        level=logging.INFO,
        format="%(asctime)s - %(levelname)s - %(message)s",
        handlers=[logging.FileHandler(path, encoding="utf-8"),
                  logging.StreamHandler(sys.stdout)],
        force=True,
    )
    return path


def ensure_folders() -> None:
    for folder in (INPUT_DIR, OUTPUT_DIR, ARCHIVE_DIR, LOG_DIR):
        folder.mkdir(parents=True, exist_ok=True)


def validate_required_files() -> dict[str, Path]:
    files = {
        "schedule": INPUT_DIR / SCHEDULE_FILENAME,
        "pharmacy": INPUT_DIR / PHARMACY_FILENAME,
        "pdf": INPUT_DIR / PDF_FILENAME,
    }
    missing = [p for p in files.values() if not p.exists()]
    if missing:
        raise FileNotFoundError("Required files are missing:\n" +
                                "\n".join(f"  - {p}" for p in missing))
    return files


def get_available_path(path: Path) -> Path:
    if not path.exists():
        return path
    return path.with_name(f"{path.stem}_{datetime.now():%Y%m%d_%H%M%S}{path.suffix}")


def archive_input_files(paths: list[Path], administration_date: date) -> Path:
    folder = (ARCHIVE_DIR / date.today().isoformat() /
              f"Administration_{administration_date.isoformat()}")
    folder.mkdir(parents=True, exist_ok=True)
    for source in paths:
        destination = get_available_path(folder / source.name)
        logging.info("Archiving %s to %s", source, destination)
        shutil.move(str(source), str(destination))
    return folder


def parse_schedule(path: Path) -> list[ScheduleRow]:
    """Parse the current patient-grouped CUSTOM schedule export."""
    workbook = xlrd.open_workbook(path)
    sheet = workbook.sheet_by_index(0)
    sample = " ".join(clean(sheet.cell_value(r, c))
                      for r in range(min(20, sheet.nrows))
                      for c in range(sheet.ncols))
    if SCHEDULE_REPORT_TITLE.casefold() not in sample.casefold():
        raise ValueError(f"Unexpected schedule report format: {path.name}")

    output: list[ScheduleRow] = []
    current_patient = ""
    current_nhs = ""
    current_provider = ""
    current_regimen = ""
    current_comment = ""

    for row_index in range(sheet.nrows):
        row = sheet.row_values(row_index) + [""] * 20
        patient = clean(row[1])
        dob = parse_excel_datetime(row[9], workbook.datemode)
        nhs = normalise_nhs_number(row[10])

        if patient and dob and nhs:
            current_patient = patient
            current_nhs = nhs
            current_provider = clean(row[18])
            current_regimen = ""
            current_comment = ""
            continue

        time_value = row[9]
        event_dt = parse_excel_datetime(time_value, workbook.datemode)
        event = clean(row[10])
        time_text = clean(time_value)

        if current_patient and not event_dt and time_text.casefold().startswith("plan "):
            current_regimen = time_text
            if clean(row[18]):
                current_comment = clean(row[18])
            continue

        if current_patient and event_dt and event:
            output.append(ScheduleRow(
                nhs_number=current_nhs,
                patient=current_patient,
                event_dt=event_dt,
                event=clean(" | ".join(
                    x for x in (current_regimen, event.lstrip("*").strip()) if x
                )),
                visit_provider=current_provider,
                location=clean(row[15]),
                comments=current_comment,
                source_row=row_index + 1,
            ))
            continue

        row_comment = clean(row[18])
        if (current_patient and row_comment
                and "report name:" not in row_comment.casefold()):
            current_comment = clean(" | ".join(dict.fromkeys(
                x for x in (current_comment, row_comment) if x
            )))

    if not output:
        raise ValueError(f"No schedule event records found in {path.name}")
    return output


def parse_pharmacy_requirements(path: Path) -> list[PharmacyRequirement]:
    workbook = xlrd.open_workbook(path)
    sheet = workbook.sheet_by_index(0)
    sample = " ".join(clean(sheet.cell_value(r, c))
                      for r in range(min(10, sheet.nrows))
                      for c in range(sheet.ncols))
    if PHARMACY_REPORT_TITLE.casefold() not in sample.casefold():
        raise ValueError(f"Unexpected Pharmacy Requirements report: {path.name}")
    output = []
    for row_index in range(sheet.nrows):
        row = sheet.row_values(row_index) + [""] * 11
        agent, patient = clean(row[0]), clean(row[4])
        administration_date = parse_text_date(row[7])
        if not (agent and patient and administration_date):
            continue
        if agent.casefold().startswith("total ") or "oncology day unit" in agent.casefold():
            continue
        output.append(PharmacyRequirement(
            administration_date, agent, format_dose(row[1], row[2]), patient,
            clean(row[6]), clean(row[9]), clean(row[10]), row_index + 1,
        ))
    if not output:
        raise ValueError(f"No medication records found in {path.name}")
    seen, distinct = set(), []
    for item in output:
        key = (item.administration_date, patient_match_key(item.patient),
               normalise(item.agent), normalise(item.dose), normalise(item.physician))
        if key not in seen:
            seen.add(key)
            distinct.append(item)
    return distinct


def likely_pdf_agent_start(line: str) -> tuple[str, str] | None:
    if not line.strip() or PDF_ADMIN_RE.match(line) or PDF_FOOTER_RE.match(line):
        return None
    match = re.match(
        r"^\s*(?P<agent>[A-Za-z0-9][A-Za-z0-9 /&().,'+\-]{1,70}?)"
        r"\s{2,}(?P<course>.+?)\s*$", line,
    )
    if not match:
        return None
    agent = clean(match.group("agent"))
    course = PDF_DATE_AT_END_RE.sub("", clean(match.group("course")))
    if agent.casefold() in {"active", "agent", "nhs number", "allergies"}:
        return None
    if not re.search(r"\b(mg|mcg|g|mmol|mL|IU|Units|dose\(s\)|tablet|capsule|"
                     r"injection|infusion|prescribed:)\b", course, re.I):
        return None
    return agent, course


def parse_pdf(path: Path) -> tuple[list[PdfPatient], list[PdfMedication]]:
    """Return every patient found plus any Chemo lines that can be parsed."""
    reader = PdfReader(str(path))
    if not reader.pages:
        raise ValueError("The Patient Medications PDF contains no pages")
    first = reader.pages[0].extract_text(extraction_mode="layout") or ""
    if PDF_REPORT_TITLE.casefold() not in first.casefold():
        raise ValueError(f"Unexpected Patient Medications report: {path.name}")

    patients, medications = [], []
    for page_number, page in enumerate(reader.pages, start=1):
        text = page.extract_text(extraction_mode="layout") or ""
        match = PDF_PATIENT_RE.search(text)
        if not match:
            continue
        visit_date = datetime.strptime(match.group("date"), "%b %d, %Y").date()
        visit_time = datetime.combine(
            visit_date, datetime.strptime(match.group("time"), "%H:%M").time())
        patient = clean(match.group("patient"))
        nhs = normalise_nhs_number(match.group("nhs"))
        patients.append(PdfPatient(visit_date, visit_time, nhs, patient, page_number))

        current_section = ""
        current = None

        def emit() -> None:
            nonlocal current
            if current and normalise(current["section"]) == "chemo":
                description = clean(" ".join(current["parts"]))
                medications.append(PdfMedication(
                    visit_date, visit_time, nhs, patient, current["section"],
                    current["agent"], description, extract_pdf_dose(description),
                    extract_route(description), page_number,
                ))
            current = None

        for raw_line in text.splitlines():
            line, stripped = raw_line.rstrip(), clean(raw_line)
            section_match = PDF_SECTION_RE.match(stripped)
            if section_match:
                emit()
                current_section = section_match.group(1)
                continue
            if PDF_FOOTER_RE.match(stripped):
                emit()
                break
            if PDF_ADMIN_RE.match(stripped):
                emit()
                continue
            start = likely_pdf_agent_start(line)
            if start:
                emit()
                current = {"section": current_section, "agent": start[0], "parts": [start[1]]}
            elif current and stripped:
                current["parts"].append(PDF_DATE_AT_END_RE.sub("", stripped))
        emit()

    patient_seen, distinct_patients = set(), []
    for item in patients:
        key = (item.visit_date, item.nhs_number)
        if key not in patient_seen:
            patient_seen.add(key)
            distinct_patients.append(item)
    med_seen, distinct_meds = set(), []
    for item in medications:
        key = (item.visit_date, item.nhs_number, normalise(item.agent),
               normalise(item.course_description))
        if key not in med_seen:
            med_seen.add(key)
            distinct_meds.append(item)
    return distinct_patients, distinct_meds


def schedule_context(appointments: list[ScheduleRow]) -> tuple:
    ordered = sorted(appointments, key=lambda x: x.event_dt)
    return (
        ordered[0].event_dt.time(),
        " | ".join(dict.fromkeys(x.event for x in ordered if x.event)),
        " | ".join(dict.fromkeys(x.visit_provider for x in ordered if x.visit_provider)),
    )


def output_row(administration_date, appointment_time, patient_id, patient,
               physician, provider, events, drug, dose, route, dispensed,
               verified, source, reference, status) -> list:
    return [administration_date, appointment_time, patient_id, patient, physician,
            provider, events, drug, dose, route, dispensed, verified, source,
            reference, status]


def add_table(ws, name: str) -> None:
    if ws.max_row < 2:
        return
    ref = f"A1:{ws.cell(ws.max_row, ws.max_column).coordinate}"
    table = Table(displayName=name, ref=ref)
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    ws.add_table(table)


def style_output_sheet(ws) -> None:
    ws.freeze_panes = "A2"
    ws.sheet_view.showGridLines = False
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[1].height = 32
    widths = {"A": 18, "B": 14, "C": 15, "D": 31, "E": 35, "F": 35,
              "G": 48, "H": 38, "I": 22, "J": 18, "K": 12, "L": 12,
              "M": 26, "N": 18, "O": 56}
    for col, width in widths.items():
        ws.column_dimensions[col].width = width
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    for cell in ws["A"][1:]:
        cell.number_format = "dd/mm/yyyy"
    for cell in ws["B"][1:]:
        cell.number_format = "hh:mm"


def write_output_sheet(wb, title: str, rows: list[list], table_name: str):
    ws = wb.create_sheet(title)
    ws.append(OUTPUT_HEADERS)
    for row in rows:
        ws.append(row)
    style_output_sheet(ws)
    add_table(ws, table_name)
    return ws


def build_workbook(schedule_path: Path, pharmacy_path: Path, pdf_path: Path,
                   output_path: Path) -> dict:
    schedule = parse_schedule(schedule_path)
    pharmacy = parse_pharmacy_requirements(pharmacy_path)
    pdf_patients, pdf_medications = parse_pdf(pdf_path)

    schedule_dates = {x.event_dt.date() for x in schedule}
    pharmacy_dates = {x.administration_date for x in pharmacy}
    pdf_dates = {x.visit_date for x in pdf_patients}
    if len(schedule_dates) != 1:
        raise ValueError("The schedule must cover exactly one administration date")
    administration_date = next(iter(schedule_dates))
    if pharmacy_dates != {administration_date}:
        raise ValueError(f"Pharmacy Requirements date(s) {sorted(pharmacy_dates)} do not match {administration_date}")
    if pdf_dates and pdf_dates != {administration_date}:
        raise ValueError(f"Patient Medications date(s) {sorted(pdf_dates)} do not match {administration_date}")

    schedule_by_nhs, schedule_by_name = defaultdict(list), defaultdict(list)
    for item in schedule:
        schedule_by_nhs[(item.event_dt.date(), item.nhs_number)].append(item)
        schedule_by_name[(item.event_dt.date(), patient_match_key(item.patient))].append(item)

    pharmacy_by_name = defaultdict(list)
    for item in pharmacy:
        pharmacy_by_name[(item.administration_date, patient_match_key(item.patient))].append(item)
    pdf_patients_by_key = {(x.visit_date, x.nhs_number): x for x in pdf_patients}
    pdf_meds_by_key = defaultdict(list)
    for item in pdf_medications:
        pdf_meds_by_key[(item.visit_date, item.nhs_number)].append(item)

    pharmacy_list, patient_review, drug_review = [], [], []

    # Walk the schedule first so Pharmacy List and Patient Review are time ordered.
    schedule_patient_keys = sorted(
        schedule_by_nhs,
        key=lambda k: min(x.event_dt for x in schedule_by_nhs[k]),
    )
    for key in schedule_patient_keys:
        appointments = schedule_by_nhs[key]
        if key[1].casefold() == "reserved":
            continue
        appointment_time, events, provider = schedule_context(appointments)
        patient = appointments[0].patient
        pharm_lines = pharmacy_by_name.get((key[0], patient_match_key(patient)), [])
        pdf_patient = pdf_patients_by_key.get(key)

        if pharm_lines:
            # Preferred source: do not duplicate the same patient from the PDF.
            for med in pharm_lines:
                pharmacy_list.append(output_row(
                    key[0], appointment_time, key[1], patient, med.physician,
                    provider, events, med.agent, med.dose, "", med.dispensed,
                    med.verified, "Pharmacy Requirements XLS", f"Row {med.source_row}",
                    "Approved",
                ))
        elif pdf_patient:
            meds = pdf_meds_by_key.get(key, [])
            if meds:
                for med in meds:
                    reason = "Planned" if med.agent or med.dose else "Planned; drug and dose not parsed"
                    pharmacy_list.append(output_row(
                        key[0], appointment_time, key[1], patient, "", provider,
                        events, med.agent, med.dose, med.route, "", "",
                        "Patient Medications PDF", f"Page {med.page}", reason,
                    ))
            else:
                # Presence in ptmeds is sufficient for Pharmacy List even when
                # no medication line can be extracted from the page.
                pharmacy_list.append(output_row(
                    key[0], appointment_time, key[1], patient, "", provider,
                    events, "", "", "", "", "", "Patient Medications PDF",
                    f"Page {pdf_patient.page}", "Planned; drug and dose not parsed",
                ))
        else:
            patient_review.append(output_row(
                key[0], appointment_time, key[1], patient, "", provider, events,
                "", "", "", "", "", "Schedule XLS",
                " | ".join(f"Row {x.source_row}" for x in appointments),
                "Not found in Pharmacy Requirements or Patient Medications",
            ))

    # Pharmacy lines not present in the schedule.
    for med in pharmacy:
        appointments = schedule_by_name.get(
            (med.administration_date, patient_match_key(med.patient)), [])
        if appointments:
            continue
        drug_review.append(output_row(
            med.administration_date, None, "", med.patient, med.physician, "", "",
            med.agent, med.dose, "", med.dispensed, med.verified,
            "Pharmacy Requirements XLS", f"Row {med.source_row}",
            "Approved medication patient not found in schedule",
        ))

    # PDF patients not present in the schedule. If nothing parsed, retain a blank line.
    for key, pdf_patient in pdf_patients_by_key.items():
        if key in schedule_by_nhs:
            continue
        meds = pdf_meds_by_key.get(key, [])
        if meds:
            for med in meds:
                drug_review.append(output_row(
                    med.visit_date, med.visit_time.time(), med.nhs_number, med.patient,
                    "", "", "", med.agent, med.dose, med.route, "", "",
                    "Patient Medications PDF", f"Page {med.page}",
                    "Planned medication patient not found in schedule",
                ))
        else:
            drug_review.append(output_row(
                pdf_patient.visit_date, pdf_patient.visit_time.time(),
                pdf_patient.nhs_number, pdf_patient.patient, "", "", "", "", "",
                "", "", "", "Patient Medications PDF", f"Page {pdf_patient.page}",
                "Patient not found in schedule; drug and dose not parsed",
            ))

    def sort_key(row):
        return (row[0], row[1] or datetime.max.time(), clean(row[3]), clean(row[7]))
    for rows in (pharmacy_list, patient_review, drug_review):
        rows.sort(key=sort_key)

    wb = Workbook()
    wb.remove(wb.active)
    write_output_sheet(wb, "Pharmacy List", pharmacy_list, "PharmacyList")
    write_output_sheet(wb, "Patient Review", patient_review, "PatientReview")
    write_output_sheet(wb, "Drug Review", drug_review, "DrugReview")

    summary = wb.create_sheet("Summary Sheet")
    summary.append(["Item", "Value"])
    summary_rows = [
        ("Generated", datetime.now()),
        ("Administration date", administration_date),
        ("Schedule source", schedule_path.name),
        ("Pharmacy source", pharmacy_path.name),
        ("Patient Medications source", pdf_path.name),
        ("Schedule event rows", len(schedule)),
        ("Distinct scheduled patients", len(schedule_patient_keys)),
        ("Distinct Pharmacy Requirements lines", len(pharmacy)),
        ("Patient Medications patients", len(pdf_patients)),
        ("Parsed Patient Medications drug lines", len(pdf_medications)),
        ("Pharmacy List lines", len(pharmacy_list)),
        ("Patient Review patients", len(patient_review)),
        ("Drug Review lines", len(drug_review)),
        ("Status meaning", "Approved = sourced from Pharmacy Requirements; Planned = sourced from Patient Medications."),
        ("Important", "Planning extract only. Verify patient, drug and dose against the current authorised ARIA prescription before preparation, release or administration."),
    ]
    for row in summary_rows:
        summary.append(row)
    summary.freeze_panes = "A2"
    summary.sheet_view.showGridLines = False
    summary.column_dimensions["A"].width = 48
    summary.column_dimensions["B"].width = 110
    for cell in summary[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
    for row in summary.iter_rows():
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    summary["B2"].number_format = "dd/mm/yyyy hh:mm"
    summary["B3"].number_format = "dd/mm/yyyy"
    add_table(summary, "SummarySheet")

    wb.save(output_path)
    return {
        "administration_date": administration_date,
        "pharmacy_list": len(pharmacy_list),
        "patient_review": len(patient_review),
        "drug_review": len(drug_review),
    }


def main() -> int:
    log_path = configure_logging()
    logging.info("Starting pharmacy planning process")
    try:
        ensure_folders()
        files = validate_required_files()
        schedule = parse_schedule(files["schedule"])
        dates = {x.event_dt.date() for x in schedule}
        if len(dates) != 1:
            raise ValueError("The schedule must cover exactly one administration date")
        administration_date = next(iter(dates))
        requested = OUTPUT_DIR / f"Pharmacy_Advance_Preparation_{administration_date}.xlsx"
        output_path = get_available_path(requested)
        stats = build_workbook(files["schedule"], files["pharmacy"], files["pdf"], output_path)
        if not output_path.exists() or output_path.stat().st_size == 0:
            raise RuntimeError("The output workbook was not created correctly")
        archive_folder = archive_input_files(
            [files["schedule"], files["pharmacy"], files["pdf"]], administration_date)
        logging.info("Pharmacy List lines: %s", stats["pharmacy_list"])
        logging.info("Patient Review patients: %s", stats["patient_review"])
        logging.info("Drug Review lines: %s", stats["drug_review"])
        print("\nPharmacy planning process completed successfully.")
        print(f"Output:  {output_path}")
        print(f"Archive: {archive_folder}")
        print(f"Log:     {log_path}\n")
        return 0
    except Exception:
        logging.exception("Pharmacy planning process failed")
        print("\nThe pharmacy planning process failed.")
        print("The input files have not been archived.")
        print(f"Review the log file: {log_path}\n")
        return 1


if __name__ == "__main__":
    sys.exit(main())
